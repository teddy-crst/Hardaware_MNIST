from playwright.sync_api import sync_playwright
from pathlib import Path
from urllib.parse import urlparse, unquote
from email.message import Message

import csv
import re
import time
import mimetypes

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from pypdf import PdfReader
from docx import Document


# ============================================================
# CONFIGURATION
# ============================================================

URL_COURS = "https://ametice.univ-amu.fr/course/view.php?id=156458"

# Sections du cours à parcourir (expressions régulières, sans
# tenir compte de la casse). Le titre "Stage de M2" ne suffit
# pas : il apparaît déjà dans la section "Généralités", qui ne
# contient que le forum.
#
# Par défaut : les sujets de l'année et toutes les archives de
# stages. Pour ne prendre que l'année en cours :
# SECTIONS_INCLUDE = r"sujets stage m2"
SECTIONS_INCLUDE = r"sujet|archive|offre"
SECTIONS_EXCLUDE = r"g[ée]n[ée]ralit|directive|d[ée]p[oô]t|soutenance|dumas|th[èe]se"

PROFILE_DIR = Path("ametice_profile")
DOWNLOAD_DIR = Path("documents_stages")
OUTPUT_CSV = Path("stages_ametice.csv")
OUTPUT_XLSX = Path("stages_ametice.xlsx")

HEADLESS = False

# Temps entre deux fiches pour ne pas marteler AMeTICE
PAUSE_BETWEEN_STAGES = 0.6

# Si le CSV existe déjà, ne retraiter que les fiches absentes
# ou en erreur (utile après une coupure).
REPRENDRE = True


# ============================================================
# OUTILS TEXTE
# ============================================================

def clean_text(text):
    if not text:
        return ""
    return re.sub(r"\s+", " ", text).strip()


def safe_filename(text, max_length=100):
    text = unquote(text or "")
    text = re.sub(r'[<>:"/\\|?*]', "_", text)
    text = re.sub(r"\s+", "_", text).strip("._ ")
    return text[:max_length] or "document"


def unique_path(path):
    if not path.exists():
        return path

    stem = path.stem
    suffix = path.suffix

    i = 2
    while True:
        candidate = path.with_name(f"{stem}_{i}{suffix}")
        if not candidate.exists():
            return candidate
        i += 1


def module_type(url):
    if "/mod/resource/" in url:
        return "resource"

    if "/mod/page/" in url:
        return "page"

    if "/mod/url/" in url:
        return "url"

    if "/mod/folder/" in url:
        return "folder"

    return "other"


def with_redirect(url):
    # Sans redirect=1, Moodle renvoie souvent une page HTML
    # (affichage intégré) au lieu du fichier ou du lien externe.
    if "/mod/resource/view.php" in url or "/mod/url/view.php" in url:
        if "redirect=1" not in url:
            return url + ("&" if "?" in url else "?") + "redirect=1"

    return url


def section_wanted(title):
    if not title:
        return False

    if SECTIONS_EXCLUDE and re.search(SECTIONS_EXCLUDE, title, re.I):
        return False

    return bool(re.search(SECTIONS_INCLUDE, title, re.I))


# ============================================================
# EXTRACTION DE TEXTE DES DOCUMENTS
# ============================================================

def extract_pdf_text(path):
    try:
        reader = PdfReader(str(path))

        texts = []

        for page in reader.pages:
            text = page.extract_text()
            if text:
                texts.append(text)

        return "\n".join(texts).strip()

    except Exception as exc:
        return f"[Erreur lecture PDF : {exc}]"


def extract_docx_text(path):
    try:
        document = Document(str(path))

        texts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                texts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                texts.append(
                    " | ".join(
                        clean_text(cell.text)
                        for cell in row.cells
                    )
                )

        return "\n".join(texts).strip()

    except Exception as exc:
        return f"[Erreur lecture DOCX : {exc}]"


def extract_text_from_file(path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf_text(path)

    if suffix == ".docx":
        return extract_docx_text(path)

    if suffix in [".txt", ".csv"]:
        try:
            return path.read_text(
                encoding="utf-8",
                errors="ignore"
            )
        except Exception:
            return ""

    return ""


# ============================================================
# NOM DU FICHIER HTTP
# ============================================================

def filename_from_response(response, fallback_name):
    headers = response.headers

    content_disposition = headers.get(
        "content-disposition",
        ""
    )

    if content_disposition:
        msg = Message()
        msg["content-disposition"] = content_disposition

        filename = msg.get_filename()

        if filename:
            return safe_filename(filename, 180)

    parsed = urlparse(response.url)

    url_name = Path(
        unquote(parsed.path)
    ).name

    if url_name and "." in url_name:
        return safe_filename(url_name, 180)

    content_type = headers.get(
        "content-type",
        ""
    ).split(";")[0]

    extension = mimetypes.guess_extension(
        content_type
    ) or ""

    return safe_filename(
        fallback_name
    ) + extension


# ============================================================
# TÉLÉCHARGEMENT VIA SESSION AUTHENTIFIÉE PLAYWRIGHT
# ============================================================

def download_file(context, url, fallback_name):
    try:
        response = context.request.get(
            with_redirect(url),
            timeout=60000,
            fail_on_status_code=False
        )

        if not response.ok:
            return None, "", (
                f"HTTP {response.status}"
            )

        content_type = response.headers.get(
            "content-type",
            ""
        ).lower()

        # Si Moodle renvoie simplement une page HTML,
        # ce n'est probablement pas encore le fichier.
        if "text/html" in content_type:
            return None, "", "Réponse HTML"

        filename = filename_from_response(
            response,
            fallback_name
        )

        path = unique_path(
            DOWNLOAD_DIR / filename
        )

        path.write_bytes(
            response.body()
        )

        text = extract_text_from_file(path)

        return path, text, ""

    except Exception as exc:
        return None, "", str(exc)


# ============================================================
# TEXTE PRINCIPAL D'UNE PAGE
# ============================================================

def get_page_text(page):
    selectors = [
        "#region-main",
        "div[role='main']",
        "main",
        ".course-content",
        "body"
    ]

    for selector in selectors:
        locator = page.locator(selector)

        if locator.count() > 0:
            try:
                text = locator.first.inner_text(
                    timeout=3000
                )

                if clean_text(text):
                    return text.strip()

            except Exception:
                pass

    return ""


def get_page_title(page, fallback):
    selectors = [
        "h1",
        ".page-header-headings h1",
        "h2"
    ]

    for selector in selectors:
        locator = page.locator(selector)

        if locator.count():
            try:
                value = clean_text(
                    locator.first.inner_text()
                )

                if value:
                    return value

            except Exception:
                pass

    try:
        value = clean_text(page.title())

        if value:
            return value

    except Exception:
        pass

    return fallback


# ============================================================
# RÉCUPÉRER LES FICHIERS D'UNE PAGE/FOLDER
# ============================================================

def extract_download_links(page):
    # Uniquement le contenu de la fiche : les liens "activité
    # précédente / suivante" de Moodle pointent vers d'autres
    # ressources du cours et feraient télécharger les voisines.
    links = page.eval_on_selector_all(
        "#region-main a[href], div[role='main'] a[href]",
        """
        els => els
            .filter(a => !a.closest(
                '.activity-navigation, #prev-activity-link, '
                + '#next-activity-link, .breadcrumb, nav'
            ))
            .map(a => ({
                text: (a.innerText || '').trim(),
                href: a.href
            }))
        """
    )

    result = []
    seen = set()

    for link in links:
        href = link.get("href", "")

        if not href:
            continue

        interesting = (
            "pluginfile.php" in href
            or "/mod/resource/view.php" in href
        )

        if not interesting:
            continue

        if href in seen:
            continue

        seen.add(href)

        result.append({
            "text": clean_text(
                link.get("text", "")
            ),
            "href": href
        })

    return result


# ============================================================
# EXTRACTION DES LIENS DE STAGES DEPUIS LA PAGE DU COURS
# ============================================================

# Nom de l'activité sans le suffixe caché "Fichier", "URL",
# "Page" ou "Dossier" que Moodle ajoute pour les lecteurs d'écran.
ACTIVITY_LINKS_JS = """
els => els.map(a => {
    const inst = a.querySelector('.instancename');
    let name = '';

    if (inst) {
        const copy = inst.cloneNode(true);
        copy.querySelectorAll('.accesshide').forEach(x => x.remove());
        name = copy.textContent;
    } else {
        name = a.innerText || a.textContent || '';
    }

    return { name: name.trim(), href: a.href };
})
"""


def extract_stage_links(page):
    allowed_patterns = [
        "/mod/page/view.php",
        "/mod/resource/view.php",
        "/mod/url/view.php",
        "/mod/folder/view.php"
    ]

    section_selectors = [
        "li.section",
        ".course-section"
    ]

    sections = None

    for selector in section_selectors:
        locator = page.locator(selector)

        if locator.count():
            sections = locator
            break

    if sections is None:
        print(
            "Sections Moodle introuvables : "
            "analyse de toute la page."
        )

    stages = []
    seen = set()

    count = sections.count() if sections is not None else 1

    for i in range(count):

        if sections is not None:
            section = sections.nth(i)

            title = ""
            title_locator = section.locator(
                ".sectionname, h3"
            )

            if title_locator.count():
                try:
                    title = clean_text(
                        title_locator.first.inner_text()
                    )
                except Exception:
                    pass

            if not section_wanted(title):
                print(f"  section ignorée : {title or '(sans titre)'}")
                continue

            print(f"  section retenue : {title}")

        else:
            section = page.locator(
                "#region-main, main, body"
            ).first
            title = ""

        # Un Locator n'a pas eval_on_selector_all (réservé à Page) :
        # evaluate_all fait la même chose sur ses liens.
        links = section.locator(
            "a[href]"
        ).evaluate_all(
            ACTIVITY_LINKS_JS
        )

        for link in links:
            name = clean_text(
                link.get("name", "")
            )

            href = link.get(
                "href",
                ""
            )

            if not name or not href:
                continue

            if not any(
                pattern in href
                for pattern in allowed_patterns
            ):
                continue

            # Moodle peut afficher plusieurs <a> pour une même activité
            if href in seen:
                continue

            seen.add(href)

            stages.append({
                "section": title,
                "nom": name,
                "url": href,
                "type": module_type(href)
            })

    return stages


# ============================================================
# ANALYSE D'UNE FICHE DE STAGE
# ============================================================

def analyse_stage(context, stage):
    result = {
        "section": stage.get("section", ""),
        "nom_liste": stage["nom"],
        "type": stage["type"],
        "url_ametice": stage["url"],
        "url_finale": "",
        "titre": "",
        "details": "",
        "documents": "",
        "erreur": ""
    }

    stage_type = stage["type"]

    # --------------------------------------------------------
    # RESOURCE = fichier direct Moodle
    # --------------------------------------------------------

    if stage_type == "resource":

        path, text, error = download_file(
            context,
            stage["url"],
            stage["nom"]
        )

        result["titre"] = stage["nom"]
        result["url_finale"] = stage["url"]

        if path:
            result["documents"] = str(path)
            result["details"] = text
        else:
            result["erreur"] = error

        return result

    # --------------------------------------------------------
    # PAGE / URL / FOLDER
    # --------------------------------------------------------

    detail_page = context.new_page()

    try:
        response = detail_page.goto(
            with_redirect(stage["url"]),
            wait_until="domcontentloaded",
            timeout=60000
        )

        detail_page.wait_for_timeout(800)

        result["url_finale"] = detail_page.url

        result["titre"] = get_page_title(
            detail_page,
            stage["nom"]
        )

        page_text = get_page_text(
            detail_page
        )

        result["details"] = page_text

        # ----------------------------------------------------
        # Vérifier si l'URL finale est directement un fichier
        # ----------------------------------------------------

        content_type = ""

        if response:
            try:
                content_type = (
                    response.headers.get(
                        "content-type",
                        ""
                    )
                    .lower()
                )
            except Exception:
                pass

        if (
            "application/pdf" in content_type
            or "application/vnd" in content_type
            or "application/octet-stream" in content_type
        ):

            path, text, error = download_file(
                context,
                detail_page.url,
                stage["nom"]
            )

            if path:
                result["documents"] = str(path)

                if text:
                    result["details"] = text

            elif error:
                result["erreur"] = error

        # ----------------------------------------------------
        # Rechercher aussi les pièces jointes
        # ----------------------------------------------------

        file_links = extract_download_links(
            detail_page
        )

        downloaded = []
        extra_texts = []

        for j, link in enumerate(
            file_links,
            start=1
        ):
            fallback = (
                link["text"]
                or f"{stage['nom']}_{j}"
            )

            path, text, error = download_file(
                context,
                link["href"],
                fallback
            )

            if path:
                downloaded.append(
                    str(path)
                )

                if text:
                    extra_texts.append(text)

        if downloaded:

            existing = result[
                "documents"
            ]

            all_docs = []

            if existing:
                all_docs.append(existing)

            all_docs.extend(downloaded)

            result["documents"] = " | ".join(
                all_docs
            )

        # Si le PDF/DOCX contient plus d'informations que la page
        if extra_texts:

            result["details"] = (
                result["details"]
                + "\n\n"
                + "\n\n".join(extra_texts)
            ).strip()

    except Exception as exc:
        result["erreur"] = str(exc)

    finally:
        detail_page.close()

    return result


# ============================================================
# EXPORT CSV
# ============================================================

FIELDS = [
    "section",
    "nom_liste",
    "type",
    "titre",
    "url_ametice",
    "url_finale",
    "details",
    "documents",
    "erreur"
]


def export_csv(results):
    with OUTPUT_CSV.open(
        "w",
        encoding="utf-8-sig",
        newline=""
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=FIELDS,
            delimiter=";"
        )

        writer.writeheader()
        writer.writerows(results)


def load_previous():
    if not (REPRENDRE and OUTPUT_CSV.exists()):
        return {}

    with OUTPUT_CSV.open(
        encoding="utf-8-sig",
        newline=""
    ) as handle:

        rows = list(
            csv.DictReader(handle, delimiter=";")
        )

    return {
        row["url_ametice"]: {field: row.get(field, "") for field in FIELDS}
        for row in rows
        if row.get("url_ametice") and not row.get("erreur")
    }


# ============================================================
# EXPORT EXCEL
# ============================================================

def export_excel(results):
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Stages"

    headers = [
        "Section",
        "Nom dans AMeTICE",
        "Type",
        "Titre",
        "URL AMeTICE",
        "URL finale",
        "Détails",
        "Documents",
        "Erreur"
    ]

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for item in results:

        sheet.append([
            item[field]
            for field in FIELDS
        ])

    widths = {
        "A": 30,
        "B": 45,
        "C": 12,
        "D": 45,
        "E": 55,
        "F": 55,
        "G": 100,
        "H": 60,
        "I": 40
    }

    for column, width in widths.items():
        sheet.column_dimensions[
            column
        ].width = width

    for row in sheet.iter_rows():

        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    sheet.freeze_panes = "A2"

    workbook.save(
        OUTPUT_XLSX
    )


# ============================================================
# PROGRAMME PRINCIPAL
# ============================================================

def main():
    DOWNLOAD_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    PROFILE_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    with sync_playwright() as p:

        context = p.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            headless=HEADLESS,
            accept_downloads=True,
            viewport={
                "width": 1440,
                "height": 900
            }
        )

        if context.pages:
            page = context.pages[0]
        else:
            page = context.new_page()

        print()
        print("Ouverture d'AMeTICE...")
        print()

        page.goto(
            URL_COURS,
            wait_until="domcontentloaded",
            timeout=60000
        )

        print(
            "Si nécessaire, connecte-toi avec ton compte AMU "
            "dans la fenêtre Chromium."
        )

        print(
            "Le programme attend que les activités Moodle apparaissent..."
        )

        # Attend indéfiniment si la connexion ENT doit être faite
        page.wait_for_selector(
            "a[href*='/mod/']",
            timeout=0
        )

        page.wait_for_timeout(
            1000
        )

        stages = extract_stage_links(
            page
        )

        print()
        print(
            f"{len(stages)} fiches potentielles trouvées."
        )
        print()

        for index, stage in enumerate(
            stages,
            start=1
        ):
            print(
                f"{index:03d}. "
                f"[{stage['type']}] "
                f"{stage['nom']}"
            )

        if not stages:
            print(
                "\nAucun stage détecté. "
                "Vérifie SECTIONS_INCLUDE ou les sélecteurs Moodle."
            )

            context.close()
            return

        previous = load_previous()

        if previous:
            print(
                f"\nReprise : {len(previous)} fiches déjà extraites "
                "seront conservées telles quelles."
            )

        results = []

        print()
        print("Début de l'extraction détaillée...")
        print()

        for i, stage in enumerate(
            stages,
            start=1
        ):

            print(
                f"[{i}/{len(stages)}] "
                f"{stage['nom']}"
            )

            if stage["url"] in previous:
                result = previous[stage["url"]]
                result["section"] = stage.get("section", "")
                results.append(result)
                print("   déjà extraite")
                continue

            result = analyse_stage(
                context,
                stage
            )

            results.append(
                result
            )

            if result["erreur"]:
                print(
                    "   Attention : "
                    + result["erreur"][:150]
                )
            else:
                chars = len(
                    result["details"]
                )

                print(
                    f"   OK - {chars} caractères extraits"
                )

            # Sauvegarde intermédiaire :
            # si le script plante, on conserve ce qui est déjà fait.
            export_csv(
                results
            )

            export_excel(
                results
            )

            time.sleep(
                PAUSE_BETWEEN_STAGES
            )

        export_csv(
            results
        )

        export_excel(
            results
        )

        print()
        print("=" * 60)
        print("EXTRACTION TERMINÉE")
        print("=" * 60)

        print(
            f"Stages : {len(results)}"
        )

        print(
            f"CSV    : {OUTPUT_CSV.resolve()}"
        )

        print(
            f"Excel  : {OUTPUT_XLSX.resolve()}"
        )

        print(
            f"Docs   : {DOWNLOAD_DIR.resolve()}"
        )

        print()
        print(
            "Étape suivante : python preparer_carte.py "
            "(fichier à déposer dans la carte)"
        )

        context.close()


if __name__ == "__main__":
    main()
